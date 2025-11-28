import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from PIL import Image
import os
import io
import json

class ExcelImageExtractor:
    def __init__(self, root):
        self.root = root
        self.root.title("Excel Image Extractor with Re-insertion")
        self.root.geometry("800x650")
        
        # Variables
        self.excel_file_path = tk.StringVar()
        self.output_folder_path = tk.StringVar()
        self.status_text = tk.StringVar(value="Ready to extract images")
        
        # Size filter variables
        self.enable_size_filter = tk.BooleanVar(value=False)
        self.filter_type = tk.StringVar(value="greater_than")
        self.size_value = tk.StringVar(value="100")
        self.size_unit = tk.StringVar(value="KB")
        
        # Re-insertion variables
        self.metadata_file_path = tk.StringVar()
        self.replace_existing_images = tk.BooleanVar(value=True)
        self.overwrite_original_file = tk.BooleanVar(value=False)
        
        # Metadata file name
        self.metadata_filename = tk.StringVar(value="extraction_data.json")
        
        self.setup_ui()
    
    def setup_ui(self):
        # Create notebook for tabs
        notebook = ttk.Notebook(self.root)
        notebook.pack(fill='both', expand=True, padx=10, pady=10)
        
        # Extract Tab
        extract_frame = ttk.Frame(notebook, padding="10")
        notebook.add(extract_frame, text="Extract Images")
        
        # Re-insert Tab
        reinsert_frame = ttk.Frame(notebook, padding="10")
        notebook.add(reinsert_frame, text="Re-insert Images")
        
        self.setup_extract_tab(extract_frame)
        self.setup_reinsert_tab(reinsert_frame)
    
    def setup_extract_tab(self, parent):
        # Configure grid weights
        parent.columnconfigure(1, weight=1)
        
        # Excel file selection
        ttk.Label(parent, text="Excel File:").grid(row=0, column=0, sticky=tk.W, pady=5)
        ttk.Entry(parent, textvariable=self.excel_file_path, width=50).grid(row=0, column=1, sticky=(tk.W, tk.E), pady=5)
        ttk.Button(parent, text="Browse", command=self.browse_excel_file).grid(row=0, column=2, padx=5, pady=5)
        
        # Output folder selection
        ttk.Label(parent, text="Output Folder:").grid(row=1, column=0, sticky=tk.W, pady=5)
        ttk.Entry(parent, textvariable=self.output_folder_path, width=50).grid(row=1, column=1, sticky=(tk.W, tk.E), pady=5)
        ttk.Button(parent, text="Browse", command=self.browse_output_folder).grid(row=1, column=2, padx=5, pady=5)
        
        # Metadata filename
        ttk.Label(parent, text="Metadata Filename:").grid(row=2, column=0, sticky=tk.W, pady=5)
        metadata_frame = ttk.Frame(parent)
        metadata_frame.grid(row=2, column=1, columnspan=2, sticky=(tk.W, tk.E), pady=5)
        metadata_frame.columnconfigure(0, weight=1)
        
        ttk.Entry(metadata_frame, textvariable=self.metadata_filename, width=30).grid(row=0, column=0, sticky=(tk.W, tk.E), padx=(0, 5))
        ttk.Button(metadata_frame, text="Change", command=self.change_metadata_filename).grid(row=0, column=1)
        
        # Size filter frame
        filter_frame = ttk.LabelFrame(parent, text="Image Size Filter", padding="5")
        filter_frame.grid(row=3, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=10)
        filter_frame.columnconfigure(1, weight=1)
        
        # Enable filter checkbox
        ttk.Checkbutton(
            filter_frame, 
            text="Enable size filtering", 
            variable=self.enable_size_filter,
            command=self.toggle_filter_options
        ).grid(row=0, column=0, columnspan=3, sticky=tk.W, pady=5)
        
        # Filter options frame (initially disabled)
        self.filter_options_frame = ttk.Frame(filter_frame)
        self.filter_options_frame.grid(row=1, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=5)
        
        # Filter type
        ttk.Label(self.filter_options_frame, text="Extract images:").grid(row=0, column=0, sticky=tk.W, padx=(0, 5))
        ttk.Radiobutton(
            self.filter_options_frame, 
            text="Greater than", 
            variable=self.filter_type, 
            value="greater_than"
        ).grid(row=0, column=1, sticky=tk.W, padx=(0, 10))
        ttk.Radiobutton(
            self.filter_options_frame, 
            text="Less than", 
            variable=self.filter_type, 
            value="less_than"
        ).grid(row=0, column=2, sticky=tk.W)
        
        # Size value and unit
        ttk.Label(self.filter_options_frame, text="Size:").grid(row=1, column=0, sticky=tk.W, pady=5)
        self.size_entry = ttk.Entry(self.filter_options_frame, textvariable=self.size_value, width=10)
        self.size_entry.grid(row=1, column=1, sticky=tk.W, padx=5, pady=5)
        
        self.size_unit_combo = ttk.Combobox(
            self.filter_options_frame, 
            textvariable=self.size_unit, 
            values=["Bytes", "KB", "MB"], 
            state="readonly", 
            width=8
        )
        self.size_unit_combo.grid(row=1, column=2, sticky=tk.W, padx=5, pady=5)
        
        # Example text
        example_text = ttk.Label(
            self.filter_options_frame, 
            text="Example: 'Greater than 100 KB' will extract images larger than 100 Kilobytes",
            foreground="gray"
        )
        example_text.grid(row=2, column=0, columnspan=4, sticky=tk.W, pady=(5, 0))
        
        # Initially disable filter options
        self.toggle_filter_options()
        
        # Extract button
        ttk.Button(parent, text="Extract Images", command=self.extract_images).grid(row=4, column=0, columnspan=3, pady=20)
        
        # Status label
        ttk.Label(parent, textvariable=self.status_text).grid(row=5, column=0, columnspan=3, pady=10)
        
        # Progress bar
        self.progress = ttk.Progressbar(parent, mode='indeterminate')
        self.progress.grid(row=6, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=5)
        
        # Results text area
        ttk.Label(parent, text="Extraction Results:").grid(row=7, column=0, sticky=tk.W, pady=(20, 5))
        
        # Text widget for results
        self.results_text = tk.Text(parent, height=15, width=70)
        self.results_text.grid(row=8, column=0, columnspan=3, sticky=(tk.W, tk.E, tk.N, tk.S), pady=5)
        
        # Scrollbar for results text
        scrollbar = ttk.Scrollbar(parent, orient=tk.VERTICAL, command=self.results_text.yview)
        scrollbar.grid(row=8, column=3, sticky=(tk.N, tk.S), pady=5)
        self.results_text.configure(yscrollcommand=scrollbar.set)
        
        # Configure row weights for parent
        parent.rowconfigure(8, weight=1)
    
    def setup_reinsert_tab(self, parent):
        # Configure grid weights
        parent.columnconfigure(1, weight=1)
        
        # Metadata file selection
        ttk.Label(parent, text="Metadata File:").grid(row=0, column=0, sticky=tk.W, pady=5)
        ttk.Entry(parent, textvariable=self.metadata_file_path, width=50).grid(row=0, column=1, sticky=(tk.W, tk.E), pady=5)
        ttk.Button(parent, text="Browse", command=self.browse_metadata_file).grid(row=0, column=2, padx=5, pady=5)
        
        # Re-insertion options frame
        options_frame = ttk.LabelFrame(parent, text="Re-insertion Options", padding="5")
        options_frame.grid(row=1, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=10)
        
        # Replace existing images option
        ttk.Checkbutton(
            options_frame,
            text="Replace existing images in Excel file",
            variable=self.replace_existing_images
        ).grid(row=0, column=0, sticky=tk.W, pady=5)
        
        # Overwrite original file option
        ttk.Checkbutton(
            options_frame,
            text="Overwrite original Excel file (instead of creating new file)",
            variable=self.overwrite_original_file
        ).grid(row=1, column=0, sticky=tk.W, pady=5)
        
        # Info about metadata
        info_text = ("The metadata file contains information about the original positions of images in the Excel file. "
                    "It is automatically created during extraction.")
        info_label = ttk.Label(parent, text=info_text, wraplength=600, foreground="blue")
        info_label.grid(row=2, column=0, columnspan=3, sticky=tk.W, pady=10)
        
        # Re-insert button
        ttk.Button(parent, text="Re-insert Images", command=self.reinsert_images).grid(row=3, column=0, columnspan=3, pady=20)
        
        # Re-insert status
        self.reinsert_status_text = tk.StringVar(value="Ready to re-insert images")
        ttk.Label(parent, textvariable=self.reinsert_status_text).grid(row=4, column=0, columnspan=3, pady=10)
        
        # Re-insert progress bar
        self.reinsert_progress = ttk.Progressbar(parent, mode='indeterminate')
        self.reinsert_progress.grid(row=5, column=0, columnspan=3, sticky=(tk.W, tk.E), pady=5)
        
        # Re-insert results text area
        ttk.Label(parent, text="Re-insertion Results:").grid(row=6, column=0, sticky=tk.W, pady=(20, 5))
        
        # Text widget for re-insert results
        self.reinsert_results_text = tk.Text(parent, height=15, width=70)
        self.reinsert_results_text.grid(row=7, column=0, columnspan=3, sticky=(tk.W, tk.E, tk.N, tk.S), pady=5)
        
        # Scrollbar for re-insert results text
        reinsert_scrollbar = ttk.Scrollbar(parent, orient=tk.VERTICAL, command=self.reinsert_results_text.yview)
        reinsert_scrollbar.grid(row=7, column=3, sticky=(tk.N, tk.S), pady=5)
        self.reinsert_results_text.configure(yscrollcommand=reinsert_scrollbar.set)
        
        # Configure row weights for parent
        parent.rowconfigure(7, weight=1)
    
    def toggle_filter_options(self):
        """Enable or disable filter options based on checkbox state"""
        if self.enable_size_filter.get():
            state = "normal"
        else:
            state = "disabled"
        
        # Update all widgets in filter options frame
        for child in self.filter_options_frame.winfo_children():
            if isinstance(child, (ttk.Radiobutton, ttk.Entry, ttk.Combobox)):
                child.configure(state=state)
    
    def change_metadata_filename(self):
        """Allow user to change the metadata filename"""
        new_name = simpledialog.askstring(
            "Metadata Filename", 
            "Enter metadata filename (with .json extension):",
            initialvalue=self.metadata_filename.get()
        )
        if new_name:
            if not new_name.endswith('.json'):
                new_name += '.json'
            self.metadata_filename.set(new_name)
    
    def browse_excel_file(self):
        file_path = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
        )
        if file_path:
            self.excel_file_path.set(file_path)
            # Set default output folder to same directory as Excel file
            default_output = os.path.join(os.path.dirname(file_path), "extracted_images")
            self.output_folder_path.set(default_output)
    
    def browse_output_folder(self):
        folder_path = filedialog.askdirectory(title="Select Output Folder")
        if folder_path:
            self.output_folder_path.set(folder_path)
    
    def browse_metadata_file(self):
        file_path = filedialog.askopenfilename(
            title="Select Metadata File",
            filetypes=[("JSON files", "*.json"), ("All files", "*.*")]
        )
        if file_path:
            self.metadata_file_path.set(file_path)
    
    def convert_to_bytes(self, size_value, unit):
        """Convert size to bytes based on unit"""
        try:
            size_value = float(size_value)
            if unit == "Bytes":
                return size_value
            elif unit == "KB":
                return size_value * 1024
            elif unit == "MB":
                return size_value * 1024 * 1024
            else:
                return size_value
        except ValueError:
            return 0
    
    def meets_size_criteria(self, image_size_bytes, threshold_bytes, filter_type):
        """Check if image meets size criteria"""
        if filter_type == "greater_than":
            return image_size_bytes > threshold_bytes
        else:  # less_than
            return image_size_bytes < threshold_bytes
    
    def format_file_size(self, size_bytes):
        """Format file size in human readable format"""
        if size_bytes >= 1024 * 1024:
            return f"{size_bytes / (1024 * 1024):.2f} MB"
        elif size_bytes >= 1024:
            return f"{size_bytes / 1024:.2f} KB"
        else:
            return f"{size_bytes} Bytes"
    
    def safe_get_image_attributes(self, image):
        """Safely get image attributes without causing AttributeError and make them JSON serializable"""
        attributes = {}
        
        # Safely get common attributes and convert to JSON-serializable types
        try:
            # Convert anchor to string representation
            if hasattr(image, 'anchor'):
                attributes['anchor'] = str(image.anchor)
            else:
                attributes['anchor'] = None
        except:
            attributes['anchor'] = None
            
        try:
            attributes['_id'] = str(image._id) if hasattr(image, '_id') else None
        except:
            attributes['_id'] = None
            
        try:
            attributes['_name'] = str(image._name) if hasattr(image, '_name') else f"image_{id(image)}"
        except:
            attributes['_name'] = f"image_{id(image)}"
            
        try:
            attributes['ref'] = str(image.ref) if hasattr(image, 'ref') else None
        except:
            attributes['ref'] = None
            
        # Try to extract cell position information if available
        try:
            if hasattr(image, 'anchor') and hasattr(image.anchor, 'from_'):
                cell_anchor = image.anchor.from_
                if hasattr(cell_anchor, 'row') and hasattr(cell_anchor, 'col'):
                    attributes['cell_position'] = {
                        'row': cell_anchor.row,
                        'col': cell_anchor.col,
                        'row_offset': getattr(cell_anchor, 'rowOff', 0),
                        'col_offset': getattr(cell_anchor, 'colOff', 0)
                    }
        except:
            attributes['cell_position'] = None
            
        return attributes
    
    def extract_image_data_safely(self, image):
        """Extract image data safely with multiple fallback methods"""
        try:
            # Method 1: Use _data()
            if hasattr(image, '_data'):
                return image._data()
            
            # Method 2: Use _blob
            if hasattr(image, '_blob'):
                return image._blob
                
            # Method 3: Try to access through ref
            if hasattr(image, 'ref') and image.ref:
                # For some versions of openpyxl, image data might be in the worksheet drawings
                return None
                
        except Exception as e:
            print(f"Error extracting image data: {e}")
            
        return None
    
    def extract_images(self):
        excel_file = self.excel_file_path.get()
        output_folder = self.output_folder_path.get()
        
        if not excel_file:
            messagebox.showerror("Error", "Please select an Excel file")
            return
        
        if not output_folder:
            messagebox.showerror("Error", "Please select an output folder")
            return
        
        if not os.path.exists(excel_file):
            messagebox.showerror("Error", "Excel file does not exist")
            return
        
        # Validate size filter inputs if enabled
        if self.enable_size_filter.get():
            try:
                size_value = float(self.size_value.get())
                if size_value <= 0:
                    messagebox.showerror("Error", "Size value must be greater than 0")
                    return
            except ValueError:
                messagebox.showerror("Error", "Please enter a valid number for size")
                return
        
        # Create output folder if it doesn't exist
        os.makedirs(output_folder, exist_ok=True)
        
        self.status_text.set("Extracting images...")
        self.progress.start()
        self.results_text.delete(1.0, tk.END)
        
        try:
            # Load the Excel workbook
            workbook = openpyxl.load_workbook(excel_file)
            
            total_images = 0
            extracted_images = []
            skipped_images = []
            metadata = {
                "source_excel": excel_file,
                "extraction_folder": output_folder,
                "images": []
            }
            
            # Calculate threshold if filtering is enabled
            threshold_bytes = None
            if self.enable_size_filter.get():
                threshold_bytes = self.convert_to_bytes(self.size_value.get(), self.size_unit.get())
                filter_type = self.filter_type.get()
                
                self.results_text.insert(tk.END, f"Size filter: {filter_type.replace('_', ' ')} {self.size_value.get()} {self.size_unit.get()}\n")
                self.results_text.insert(tk.END, f"Threshold: {self.format_file_size(threshold_bytes)}\n\n")
            
            # Iterate through all sheets
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                
                # Check if sheet has drawings (images)
                if hasattr(sheet, '_images') and sheet._images:
                    self.results_text.insert(tk.END, f"Processing sheet: {sheet_name}\n")
                    
                    # Extract images from this sheet
                    for i, image in enumerate(sheet._images):
                        try:
                            # Get image data safely
                            image_data = self.extract_image_data_safely(image)
                            
                            if image_data is None:
                                self.results_text.insert(tk.END, f"  ✗ Cannot extract image {i+1}: No image data found\n")
                                continue
                                
                            image_size = len(image_data)
                            
                            # Apply size filter if enabled
                            if self.enable_size_filter.get():
                                if not self.meets_size_criteria(image_size, threshold_bytes, self.filter_type.get()):
                                    size_str = self.format_file_size(image_size)
                                    skipped_images.append((f"Image {i+1} from {sheet_name}", size_str))
                                    self.results_text.insert(tk.END, f"  ⚬ Skipped (size): image {i+1} - {size_str}\n")
                                    continue
                            
                            # Determine image format and extension
                            img_format, file_extension = self.detect_image_format(image_data)
                            
                            if img_format:
                                # Create filename
                                filename = f"sheet_{sheet_name}_image_{i+1}.{file_extension}"
                                # Remove invalid characters from filename
                                filename = "".join(c for c in filename if c.isalnum() or c in "._-")
                                
                                filepath = os.path.join(output_folder, filename)
                                
                                # Save the image
                                with open(filepath, 'wb') as f:
                                    f.write(image_data)
                                
                                # Store metadata for re-insertion (safely get attributes)
                                image_attributes = self.safe_get_image_attributes(image)
                                
                                # Create a simplified metadata entry that's JSON serializable
                                image_metadata = {
                                    "original_sheet": sheet_name,
                                    "image_file": filename,
                                    "position": image_attributes,
                                    "size": image_size,
                                    "format": img_format,
                                    "image_index": i
                                }
                                metadata["images"].append(image_metadata)
                                
                                size_str = self.format_file_size(image_size)
                                extracted_images.append((filename, sheet_name, size_str))
                                self.results_text.insert(tk.END, f"  ✓ Extracted: {filename} - {size_str}\n")
                                total_images += 1
                            else:
                                self.results_text.insert(tk.END, f"  ✗ Unknown format for image {i+1}\n")
                                
                        except Exception as e:
                            self.results_text.insert(tk.END, f"  ✗ Error extracting image {i+1}: {str(e)}\n")
                else:
                    self.results_text.insert(tk.END, f"No images found in sheet: {sheet_name}\n")
            
            # Create a custom JSON encoder to handle any non-serializable objects
            class CustomJSONEncoder(json.JSONEncoder):
                def default(self, obj):
                    try:
                        # Try to convert to string representation
                        return str(obj)
                    except:
                        # If that fails, return a placeholder
                        return f"Non-serializable object: {type(obj)}"
            
            # Save metadata file with custom encoder
            metadata_file = os.path.join(output_folder, self.metadata_filename.get())
            with open(metadata_file, 'w') as f:
                json.dump(metadata, f, indent=2, cls=CustomJSONEncoder, ensure_ascii=False)
            
            self.progress.stop()
            
            # Display summary
            self.results_text.insert(tk.END, f"\n{'='*50}\n")
            self.results_text.insert(tk.END, "EXTRACTION SUMMARY\n")
            self.results_text.insert(tk.END, f"{'='*50}\n")
            self.results_text.insert(tk.END, f"Total images extracted: {total_images}\n")
            self.results_text.insert(tk.END, f"Metadata saved to: {metadata_file}\n")
            
            if self.enable_size_filter.get():
                self.results_text.insert(tk.END, f"Images skipped by filter: {len(skipped_images)}\n")
            
            self.results_text.insert(tk.END, f"Output folder: {output_folder}\n")
            
            if total_images > 0:
                self.status_text.set(f"Successfully extracted {total_images} images")
                messagebox.showinfo("Success", 
                                  f"Extracted {total_images} images to:\n{output_folder}\n\n"
                                  f"Metadata saved for re-insertion.")
            else:
                self.status_text.set("No images found matching criteria")
                messagebox.showinfo("No Images", "No images were found matching the specified criteria")
                
        except Exception as e:
            self.progress.stop()
            self.status_text.set("Error during extraction")
            self.results_text.insert(tk.END, f"\n❌ Error: {str(e)}")
            messagebox.showerror("Error", f"Failed to extract images: {str(e)}")
    
    def reinsert_images(self):
        metadata_file = self.metadata_file_path.get()
        
        if not metadata_file:
            messagebox.showerror("Error", "Please select a metadata file")
            return
        
        if not os.path.exists(metadata_file):
            messagebox.showerror("Error", "Metadata file does not exist")
            return
        
        # Ask for destination Excel file
        excel_file = filedialog.askopenfilename(
            title="Select Excel File to Re-insert Images",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        
        if not excel_file:
            return
        
        # Ask for output file name if not overwriting
        if self.overwrite_original_file.get():
            output_file = excel_file
            if not messagebox.askyesno("Confirm Overwrite", 
                                     f"Are you sure you want to overwrite the original file?\n\n{excel_file}"):
                return
        else:
            # Ask user for output filename
            default_name = excel_file.replace('.xlsx', '_with_images.xlsx')
            output_file = filedialog.asksaveasfilename(
                title="Save Excel File As",
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
                initialfile=os.path.basename(default_name)
            )
            
            if not output_file:
                return
        
        self.reinsert_status_text.set("Re-inserting images...")
        self.reinsert_progress.start()
        self.reinsert_results_text.delete(1.0, tk.END)
        
        try:
            # Load metadata
            with open(metadata_file, 'r') as f:
                metadata = json.load(f)
            
            # Load Excel workbook
            workbook = openpyxl.load_workbook(excel_file)
            
            extraction_folder = metadata.get("extraction_folder", os.path.dirname(metadata_file))
            reinserted_count = 0
            
            self.reinsert_results_text.insert(tk.END, f"Re-inserting images to: {output_file}\n")
            self.reinsert_results_text.insert(tk.END, f"Source folder: {extraction_folder}\n")
            
            # If replacing existing images, clear all images from sheets mentioned in metadata
            if self.replace_existing_images.get():
                self.reinsert_results_text.insert(tk.END, "Replacing existing images...\n")
                for img_meta in metadata["images"]:
                    sheet_name = img_meta["original_sheet"]
                    if sheet_name in workbook.sheetnames:
                        sheet = workbook[sheet_name]
                        # Clear all images from this sheet
                        if hasattr(sheet, '_images'):
                            sheet._images.clear()
            
            self.reinsert_results_text.insert(tk.END, "\n")
            
            # Group images by sheet
            sheet_images = {}
            for img_meta in metadata["images"]:
                sheet_name = img_meta["original_sheet"]
                if sheet_name not in sheet_images:
                    sheet_images[sheet_name] = []
                sheet_images[sheet_name].append(img_meta)
            
            for sheet_name, images in sheet_images.items():
                # Check if sheet exists
                if sheet_name not in workbook.sheetnames:
                    self.reinsert_results_text.insert(tk.END, f"✗ Sheet not found: {sheet_name}\n")
                    continue
                
                sheet = workbook[sheet_name]
                self.reinsert_results_text.insert(tk.END, f"Processing sheet: {sheet_name}\n")
                
                for i, img_meta in enumerate(images):
                    try:
                        image_file = img_meta["image_file"]
                        image_path = os.path.join(extraction_folder, image_file)
                        
                        if not os.path.exists(image_path):
                            self.reinsert_results_text.insert(tk.END, f"  ✗ Image not found: {image_file}\n")
                            continue
                        
                        # Add image to sheet
                        img = XLImage(image_path)
                        
                        # Try to use original position if available
                        position_info = img_meta.get("position", {})
                        cell_position = position_info.get("cell_position")
                        
                        if cell_position and self.replace_existing_images.get():
                            # Use original position
                            try:
                                # Convert to Excel cell reference
                                from openpyxl.utils import get_column_letter
                                col_letter = get_column_letter(cell_position['col'] + 1)  # Convert 0-indexed to 1-indexed
                                cell_ref = f"{col_letter}{cell_position['row'] + 1}"
                                sheet.add_image(img, cell_ref)
                                position_info_str = f"at original position {cell_ref}"
                            except Exception as e:
                                # Fallback to sequential placement
                                row = (i * 15) + 1
                                cell_ref = f"A{row}"
                                sheet.add_image(img, cell_ref)
                                position_info_str = f"at {cell_ref} (fallback)"
                        else:
                            # Use sequential placement
                            row = (i * 15) + 1
                            cell_ref = f"A{row}"
                            sheet.add_image(img, cell_ref)
                            position_info_str = f"at {cell_ref}"
                        
                        self.reinsert_results_text.insert(tk.END, f"  ✓ Re-inserted: {image_file} {position_info_str}\n")
                        reinserted_count += 1
                        
                    except Exception as e:
                        self.reinsert_results_text.insert(tk.END, f"  ✗ Error re-inserting {image_file}: {str(e)}\n")
            
            # Save the modified workbook
            workbook.save(output_file)
            
            self.reinsert_progress.stop()
            
            # Display summary
            self.reinsert_results_text.insert(tk.END, f"\n{'='*50}\n")
            self.reinsert_results_text.insert(tk.END, "RE-INSERTION SUMMARY\n")
            self.reinsert_results_text.insert(tk.END, f"{'='*50}\n")
            self.reinsert_results_text.insert(tk.END, f"Total images re-inserted: {reinserted_count}\n")
            self.reinsert_results_text.insert(tk.END, f"Output file: {output_file}\n")
            if self.replace_existing_images.get():
                self.reinsert_results_text.insert(tk.END, "Mode: Replaced existing images\n")
            else:
                self.reinsert_results_text.insert(tk.END, "Mode: Added images (kept existing)\n")
            
            self.reinsert_status_text.set(f"Successfully re-inserted {reinserted_count} images")
            messagebox.showinfo("Success", 
                              f"Re-inserted {reinserted_count} images to:\n{output_file}")
            
        except Exception as e:
            self.reinsert_progress.stop()
            self.reinsert_status_text.set("Error during re-insertion")
            self.reinsert_results_text.insert(tk.END, f"\n❌ Error: {str(e)}")
            messagebox.showerror("Error", f"Failed to re-insert images: {str(e)}")
    
    def detect_image_format(self, image_data):
        """Detect image format from binary data"""
        try:
            # Common image format signatures
            if image_data.startswith(b'\xff\xd8\xff'):
                return 'JPEG', 'jpg'
            elif image_data.startswith(b'\x89PNG\r\n\x1a\n'):
                return 'PNG', 'png'
            elif image_data.startswith(b'GIF8'):
                return 'GIF', 'gif'
            elif image_data.startswith(b'BM'):
                return 'BMP', 'bmp'
            elif image_data.startswith(b'\x00\x00\x01\x00'):
                return 'ICO', 'ico'
            elif image_data.startswith(b'RIFF') and image_data[8:12] == b'WEBP':
                return 'WEBP', 'webp'
            else:
                # Try to open with PIL to detect format
                try:
                    image = Image.open(io.BytesIO(image_data))
                    format_name = image.format
                    if format_name:
                        return format_name, format_name.lower()
                except:
                    pass
                
                return None, None
        except:
            return None, None

def main():
    root = tk.Tk()
    app = ExcelImageExtractor(root)
    root.mainloop()

if __name__ == "__main__":
    main()